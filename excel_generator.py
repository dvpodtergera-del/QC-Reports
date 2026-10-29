import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as OpenPyXLImage
from PIL import Image, ImageOps

def create_full_qc_report(
    report_data: dict, 
    photos_by_section: dict, 
    spec_data: dict = None, 
    measured_samples: list = None
) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "QC Inspection Report"
    ws.views.sheetView[0].showGridLines = True

    # Корпоративная палитра Bizneskit (Slate Blue & Dark Graphite)
    c_brand = "1E293B"      # Dark Slate
    c_header = "334155"     # Slate 700
    c_sub = "475569"        # Slate 600
    c_border = "CBD5E1"     # Slate 300
    c_meta_bg = "F8FAFC"    # Светло-серый фон
    
    # Статусные цвета
    c_fail_bg = "FEE2E2"
    c_fail_fg = "991B1B"
    c_pass_bg = "DCFCE7"
    c_pass_fg = "166534"
    c_pend_bg = "FEF3C7"
    c_pend_fg = "92400E"

    thin_border = Border(
        left=Side(style="thin", color=c_border),
        right=Side(style="thin", color=c_border),
        top=Side(style="thin", color=c_border),
        bottom=Side(style="thin", color=c_border)
    )

    # 1. Ширина колонок листа (A-H)
    col_widths = {
        "A": 4, 
        "B": 18, 
        "C": 26, 
        "D": 14, 
        "E": 14, 
        "F": 16, 
        "G": 16, 
        "H": 18
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    # 2. Главная брендовая плашка Bizneskit
    ws.merge_cells("B2:H2")
    ws.row_dimensions[2].height = 40
    company_title = report_data.get("company_name", "BIZNESKIT").upper()
    h_cell = ws["B2"]
    h_cell.value = f"{company_title} — QUALITY CONTROL INSPECTION REPORT"
    h_cell.fill = PatternFill(start_color=c_brand, end_color=c_brand, fill_type="solid")
    h_cell.font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    h_cell.alignment = Alignment(horizontal="center", vertical="center")

    # 3. Информационный блок партии
    res_status = report_data.get("result_status", "PENDING").upper()
    meta_rows = [
        ("REPORT NO:", report_data.get("report_no", "QC-001"), "INSPECTION DATE:", str(report_data.get("inspection_date", ""))),
        ("CLIENT:", report_data.get("client_name", "-"), "INSPECTOR ID:", str(report_data.get("inspector_id", "-"))),
        ("SUPPLIER / FACTORY:", report_data.get("supplier_name", "-"), "PO / ORDER NO:", str(report_data.get("order_no", "-"))),
        ("ITEM / SKU:", f"{report_data.get('item_name', '-')} / {report_data.get('sku', '-')}", "OVERALL RESULT:", res_status)
    ]

    curr_row = 4
    for r_data in meta_rows:
        ws.row_dimensions[curr_row].height = 22
        
        # Левая колонка
        c1 = ws.cell(row=curr_row, column=2, value=r_data[0])
        c1.fill = PatternFill(start_color=c_meta_bg, fill_type="solid")
        c1.font = Font(name="Calibri", size=9, bold=True, color=c_sub)
        c1.border = thin_border
        
        ws.merge_cells(start_row=curr_row, start_column=3, end_row=curr_row, end_column=4)
        c2 = ws.cell(row=curr_row, column=3, value=r_data[1])
        c2.font = Font(name="Calibri", size=9, bold=True, color="0F172A")
        c2.border = thin_border

        # Правая колонка
        c3 = ws.cell(row=curr_row, column=5, value=r_data[2])
        c3.fill = PatternFill(start_color=c_meta_bg, fill_type="solid")
        c3.font = Font(name="Calibri", size=9, bold=True, color=c_sub)
        c3.border = thin_border

        ws.merge_cells(start_row=curr_row, start_column=6, end_row=curr_row, end_column=8)
        c4 = ws.cell(row=curr_row, column=6, value=r_data[3])
        c4.border = thin_border
        
        if r_data[2] == "OVERALL RESULT:":
            c4.alignment = Alignment(horizontal="center", vertical="center")
            if res_status == "PASSED":
                c4.fill = PatternFill(start_color=c_pass_bg, fill_type="solid")
                c4.font = Font(name="Calibri", size=11, bold=True, color=c_pass_fg)
            elif res_status == "REJECTED":
                c4.fill = PatternFill(start_color=c_fail_bg, fill_type="solid")
                c4.font = Font(name="Calibri", size=11, bold=True, color=c_fail_fg)
            else:
                c4.fill = PatternFill(start_color=c_pend_bg, fill_type="solid")
                c4.font = Font(name="Calibri", size=11, bold=True, color=c_pend_fg)
        else:
            c4.font = Font(name="Calibri", size=9, color="0F172A")
            c4.alignment = Alignment(horizontal="left", vertical="center")

        curr_row += 1

    curr_row += 1

    # 4. Блок размерной сетки для одежды (Spec Sheet)
    if spec_data and measured_samples:
        ws.merge_cells(start_row=curr_row, start_column=2, end_row=curr_row, end_column=8)
        ws.row_dimensions[curr_row].height = 24
        sec_cell = ws.cell(row=curr_row, column=2, value="GARMENT MEASUREMENT & TOLERANCE SPEC SHEET")
        sec_cell.fill = PatternFill(start_color=c_header, fill_type="solid")
        sec_cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        sec_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        curr_row += 1

        # Заголовки таблицы замеров
        ws.row_dimensions[curr_row].height = 20
        spec_headers = ["Code", "POM Description", "Tol (-)", "Tol (+)"]
        for idx, h_name in enumerate(spec_headers, start=2):
            c = ws.cell(row=curr_row, column=idx, value=h_name)
            c.fill = PatternFill(start_color=c_sub, fill_type="solid")
            c.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = thin_border

        sample_col = 6
        for sample in measured_samples:
            ws.cell(row=curr_row, column=sample_col, value=f"#{sample['sample_no']} ({sample['size']}) Act").font = Font(name="Calibri", size=8, bold=True, color="FFFFFF")
            ws.cell(row=curr_row, column=sample_col).fill = PatternFill(start_color=c_sub, fill_type="solid")
            ws.cell(row=curr_row, column=sample_col).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=curr_row, column=sample_col).border = thin_border

            ws.cell(row=curr_row, column=sample_col + 1, value="Diff").font = Font(name="Calibri", size=8, bold=True, color="FFFFFF")
            ws.cell(row=curr_row, column=sample_col + 1).fill = PatternFill(start_color=c_sub, fill_type="solid")
            ws.cell(row=curr_row, column=sample_col + 1).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=curr_row, column=sample_col + 1).border = thin_border
            sample_col += 2

        curr_row += 1

        # Заполнение строк спецификации
        for pom in spec_data.get("poms", []):
            ws.row_dimensions[curr_row].height = 18
            ws.cell(row=curr_row, column=2, value=pom.get("code", "")).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=curr_row, column=3, value=pom.get("description", "")).alignment = Alignment(horizontal="left", vertical="center")
            ws.cell(row=curr_row, column=4, value=f"-{pom.get('tol_minus', 1.0)}").alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=curr_row, column=5, value=f"+{pom.get('tol_plus', 1.0)}").alignment = Alignment(horizontal="center", vertical="center")

            for c_i in range(2, 6):
                ws.cell(row=curr_row, column=c_i).border = thin_border
                ws.cell(row=curr_row, column=c_i).font = Font(name="Calibri", size=9)

            s_col = 6
            for sample in measured_samples:
                item = next((x for x in sample.get("items", []) if x.get("code") == pom.get("code")), None)
                if item:
                    c_act = ws.cell(row=curr_row, column=s_col, value=item.get("actual", "-"))
                    c_diff = ws.cell(row=curr_row, column=s_col + 1, value=item.get("diff", 0))
                    
                    c_act.border = thin_border
                    c_diff.border = thin_border
                    c_act.alignment = Alignment(horizontal="center", vertical="center")
                    c_diff.alignment = Alignment(horizontal="center", vertical="center")
                    c_act.font = Font(name="Calibri", size=9)

                    if item.get("status") == "FAIL":
                        c_diff.fill = PatternFill(start_color=c_fail_bg, fill_type="solid")
                        c_diff.font = Font(name="Calibri", size=9, bold=True, color=c_fail_fg)
                    else:
                        c_diff.fill = PatternFill(start_color=c_pass_bg, fill_type="solid")
                        c_diff.font = Font(name="Calibri", size=9, bold=True, color=c_pass_fg)
                s_col += 2

            curr_row += 1
        curr_row += 1

    # 5. Секции фотоматериалов
    section_titles = {
        "products": "PRODUCT WORKMANSHIP PHOTOS",
        "packing": "PACKAGING & CARTON DETAILS",
        "labels": "LABELS & BARCODES VERIFICATION",
        "defects": "DEFECT DETAILS & LOG"
    }

    for s_key, s_label in section_titles.items():
        photos = photos_by_section.get(s_key, [])
        if not photos:
            continue

        ws.merge_cells(start_row=curr_row, start_column=2, end_row=curr_row, end_column=8)
        ws.row_dimensions[curr_row].height = 22
        t_cell = ws.cell(row=curr_row, column=2, value=s_label)
        t_cell.fill = PatternFill(start_color=c_header, fill_type="solid")
        t_cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        t_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        curr_row += 1

        # Сетка: по 2 фото в строку (B и F)
        col_slots = [('B', 2), ('F', 6)]
        for i in range(0, len(photos), 2):
            ws.row_dimensions[curr_row].height = 140
            ws.row_dimensions[curr_row + 1].height = 22

            for slot_idx, (col_letter, col_num) in enumerate(col_slots):
                p_idx = i + slot_idx
                if p_idx < len(photos):
                    p_info = photos[p_idx]
                    
                    # Обработка через Pillow: коррекция EXIF и кроп
                    img = Image.open(io.BytesIO(p_info["bytes"]))
                    img = ImageOps.exif_transpose(img)
                    img = ImageOps.fit(img, (320, 180), method=Image.Resampling.LANCZOS)
                    
                    img_buf = io.BytesIO()
                    img.save(img_buf, format="JPEG", quality=85)
                    img_buf.seek(0)

                    xl_img = OpenPyXLImage(img_buf)
                    xl_img.width = 240
                    xl_img.height = 135
                    ws.add_image(xl_img, f"{col_letter}{curr_row}")

                    # Подпись под фото
                    ws.merge_cells(start_row=curr_row + 1, start_column=col_num, end_row=curr_row + 1, end_column=col_num + 2)
                    cap_cell = ws.cell(row=curr_row + 1, column=col_num, value=p_info.get("caption", f"Photo #{p_idx + 1}"))
                    cap_cell.font = Font(name="Calibri", size=8, italic=True, color=c_sub)
                    cap_cell.alignment = Alignment(horizontal="center", vertical="center")

            curr_row += 2
        curr_row += 1

    # 6. Выводы и комментарии инспекции
    ws.merge_cells(start_row=curr_row, start_column=2, end_row=curr_row, end_column=8)
    ws.row_dimensions[curr_row].height = 22
    conc_cell = ws.cell(row=curr_row, column=2, value="INSPECTION COMMENTS & CONCLUSION")
    conc_cell.fill = PatternFill(start_color=c_header, fill_type="solid")
    conc_cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    conc_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    curr_row += 1

    ws.merge_cells(start_row=curr_row, start_column=2, end_row=curr_row + 3, end_column=8)
    for r in range(curr_row, curr_row + 4):
        ws.row_dimensions[r].height = 18
    text_cell = ws.cell(row=curr_row, column=2, value=report_data.get("comments", "Inspection carried out in accordance with standard specifications."))
    text_cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    text_cell.font = Font(name="Calibri", size=9, color="0F172A")

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
